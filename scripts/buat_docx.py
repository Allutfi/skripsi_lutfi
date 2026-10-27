#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
buat_docx.py - Mengubah naskah Markdown (.md) menjadi dokumen Word (.docx) yang
sudah berformat sesuai Buku Pedoman TA FTI ITB Widya Gama Lumajang 2026-2027.
Opsional: sekalian jadi PDF.

Contoh:
    python scripts/buat_docx.py draft/proposal.md
    python scripts/buat_docx.py draft/proposal.md --pdf
    python scripts/buat_docx.py draft/proposal.md --tahap laporan --romawi

Sintaks Markdown yang didukung:
    # BAB 1 PENDAHULUAN          judul bab (kapital, tengah, bold; mulai halaman baru)
    ## 1.1 Latar Belakang Masalah   subbab      ### 1.1.1 ...   anak subbab
    paragraf biasa               (baris berurutan digabung jadi satu paragraf)
    **tebal**  *miring*  `kode`
    - butir                      daftar bullet
    1. butir / a. butir          daftar bernomor (nomor ditulis apa adanya)
    | A | B |                    tabel (baris ke-2 '|---|---|'); garis horizontal saja
    Tabel 3.1 Judul Tabel        judul tabel: tulis DI ATAS tabel
    ![Gambar 3.1 Judul](gbr/a.png)   gambar + judul di bawahnya
    ``` ... ```                  blok kode
    <!-- pagebreak -->           paksa halaman baru
    # DAFTAR PUSTAKA             entri berikutnya diformat spasi 1 + menjorok 1,5 cm

Hasil mencakup Bab 1-3 (+ Daftar Pustaka/Lampiran). Halaman sampul, pengesahan, dan
persetujuan diunduh dari SIAKAD dan digabung manual.
"""
import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import config
from common import ensure_output_dir, setup_stdout

F = config.FORMAT
FONT = F["font_name"]
SIZE = F["font_size_pt"]
MONO = "Courier New"

ALIGN = {
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
}

INLINE = re.compile(r"(\*\*.+?\*\*|`.+?`|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*))")
IMG_RE = re.compile(r"^!\[(.*?)\]\((.*?)\)\s*$")
CAPTION_RE = re.compile(r"^(tabel|gambar)\s+\d+([.\-]\d+)?", re.IGNORECASE)
NUM_RE = re.compile(r"^(\d+[.)]|[a-z][.)])\s+(.*)$")
BULLET_RE = re.compile(r"^[-*]\s+(.*)$")


# ---------------------------------------------------------------------------
# Bantu: font
# ---------------------------------------------------------------------------
def fix_rfonts(rpr, name):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), name)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        if rf.get(qn(a)) is not None:
            del rf.attrib[qn(a)]


def style_font(style, name=FONT, size=SIZE, bold=None, italic=None):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic
    fix_rfonts(style.element.get_or_add_rPr(), name)


def run_font(run, name=FONT, size=SIZE, bold=None, italic=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    fix_rfonts(run._element.get_or_add_rPr(), name)


def set_lang(style, lang="id-ID"):
    rpr = style.element.get_or_add_rPr()
    for old in rpr.findall(qn("w:lang")):
        rpr.remove(old)
    el = OxmlElement("w:lang")
    el.set(qn("w:val"), lang)
    rpr.append(el)


def setup_styles(doc, spasi):
    normal = doc.styles["Normal"]
    style_font(normal)
    set_lang(normal)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = spasi
    for lvl in (1, 2, 3, 4):
        st = doc.styles[f"Heading {lvl}"]
        style_font(st, bold=True, italic=False)
        pf = st.paragraph_format
        pf.line_spacing = spasi
        pf.keep_with_next = True
        pf.space_before = Pt(0 if lvl == 1 else 6)
        pf.space_after = Pt(12 if lvl == 1 else 0)
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if lvl == 1 else WD_ALIGN_PARAGRAPH.LEFT
    for nama in ("List Bullet", "Header", "Footer"):
        try:
            style_font(doc.styles[nama])
        except KeyError:
            pass


# ---------------------------------------------------------------------------
# Bantu: halaman, header/footer
# ---------------------------------------------------------------------------
def add_page_field(paragraph):
    run = paragraph.add_run()
    run_font(run)
    for tipe, teks in (("begin", None), (None, "PAGE"), ("end", None)):
        if tipe:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tipe)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = teks
        run._r.append(el)


def set_page_numbering(section, romawi, mulai):
    sect = section._sectPr
    for old in sect.findall(qn("w:pgNumType")):
        sect.remove(old)
    pg = OxmlElement("w:pgNumType")
    pg.set(qn("w:fmt"), "upperRoman" if romawi else "decimal")
    if mulai:
        pg.set(qn("w:start"), "1")
    pgmar = sect.find(qn("w:pgMar"))
    if pgmar is not None:
        pgmar.addnext(pg)
    else:
        sect.append(pg)


def setup_section_basic(section):
    section.page_width = Cm(F["page_width_cm"])
    section.page_height = Cm(F["page_height_cm"])
    section.top_margin = Cm(F["margin_top_cm"])
    section.left_margin = Cm(F["margin_left_cm"])
    section.right_margin = Cm(F["margin_right_cm"])
    section.bottom_margin = Cm(F["margin_bottom_cm"])
    section.header_distance = Cm(F["header_distance_cm"])
    section.footer_distance = Cm(F["footer_distance_cm"])
    section.different_first_page_header_footer = True


def setup_header_footer(section):
    """Nomor halaman: kanan atas (1,5 cm dari tepi kanan & atas); halaman pertama
    tiap bab (judul bab): tengah bawah, 1 cm dari tepi bawah."""
    hdr = section.header
    hdr.is_linked_to_previous = False
    p = hdr.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.right_indent = Cm(-F["right_offset_cm"])
    p.paragraph_format.line_spacing = 1.0
    add_page_field(p)

    ftr = section.footer
    ftr.is_linked_to_previous = False

    fh = section.first_page_header
    fh.is_linked_to_previous = False

    ff = section.first_page_footer
    ff.is_linked_to_previous = False
    p2 = ff.paragraphs[0]
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.line_spacing = 1.0
    add_page_field(p2)


# ---------------------------------------------------------------------------
# Bantu: paragraf
# ---------------------------------------------------------------------------
def add_inline(p, text, size=SIZE, bold=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run_font(p.add_run(part[2:-2]), size=size, bold=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run_font(p.add_run(part[1:-1]), name=MONO, size=max(size - 2, 9), bold=bold)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run_font(p.add_run(part[1:-1]), size=size, italic=True, bold=bold)
        else:
            run_font(p.add_run(part), size=size, bold=bold)


class Penyusun:
    def __init__(self, doc, tahap, romawi, dasar):
        self.doc = doc
        self.cfg = config.TAHAP[tahap]
        self.spasi = self.cfg["spasi"]
        self.romawi = romawi
        self.dasar = Path(dasar)
        self.ada_isi = False
        self.seksi_pertama_siap = False
        self.mode_pustaka = False
        self.peringatan = []

    # -- seksi -------------------------------------------------------------
    def mulai_bab(self):
        if not self.seksi_pertama_siap:
            s = self.doc.sections[0]
            setup_section_basic(s)
            set_page_numbering(s, self.romawi, mulai=True)
            setup_header_footer(s)
            self.seksi_pertama_siap = True
            return
        if self.ada_isi:
            s = self.doc.add_section(WD_SECTION.NEW_PAGE)
            setup_section_basic(s)
            set_page_numbering(s, self.romawi, mulai=False)

    def pastikan_seksi(self):
        if not self.seksi_pertama_siap:
            self.mulai_bab()

    # -- elemen --------------------------------------------------------------
    def heading(self, teks, level):
        if level == 1:
            self.mulai_bab()
            teks = teks.upper()
            n = re.sub(r"\s+", " ", teks.strip().lower())
            self.mode_pustaka = n.startswith("daftar pustaka")
        else:
            self.pastikan_seksi()
            teks = teks.rstrip().rstrip(".")
        p = self.doc.add_paragraph(style=f"Heading {min(level, 4)}")
        run_font(p.add_run(teks), bold=True, italic=False)
        self.ada_isi = True

    def paragraf(self, teks):
        self.pastikan_seksi()
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = self.spasi
        if CAPTION_RE.match(teks):
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pf.line_spacing = 1.0
            pf.space_after = Pt(6)
            if teks.lower().startswith("tabel"):
                pf.keep_with_next = True
        elif re.match(r"^sumber\s*:", teks, re.IGNORECASE):
            pf.line_spacing = 1.0
            pf.space_after = Pt(6)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        elif self.mode_pustaka:
            pf.line_spacing = 1.0
            pf.left_indent = Cm(F["pustaka_hanging_cm"])
            pf.first_line_indent = Cm(-F["pustaka_hanging_cm"])
            pf.space_after = Pt(6)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        else:
            pf.first_line_indent = Cm(F["first_line_indent_cm"])
            pf.alignment = ALIGN.get(config.ALIGN_BODY, WD_ALIGN_PARAGRAPH.LEFT)
        add_inline(p, teks)
        self.ada_isi = True

    def butir_bullet(self, teks):
        self.pastikan_seksi()
        p = self.doc.add_paragraph(style="List Bullet")
        pf = p.paragraph_format
        pf.line_spacing = self.spasi
        pf.alignment = ALIGN.get(config.ALIGN_BODY, WD_ALIGN_PARAGRAPH.LEFT)
        add_inline(p, teks)
        self.ada_isi = True

    def butir_nomor(self, penanda, teks):
        self.pastikan_seksi()
        p = self.doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = self.spasi
        pf.left_indent = Cm(1.25)
        pf.first_line_indent = Cm(-0.75)
        pf.tab_stops.add_tab_stop(Cm(1.25))
        pf.alignment = ALIGN.get(config.ALIGN_BODY, WD_ALIGN_PARAGRAPH.LEFT)
        run_font(p.add_run(penanda + "\t"))
        add_inline(p, teks)
        self.ada_isi = True

    def kode(self, baris):
        self.pastikan_seksi()
        for b in baris:
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.line_spacing = 1.0
            pf.left_indent = Cm(0.75)
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run_font(p.add_run(b if b else " "), name=MONO, size=10)
        self.ada_isi = True

    def page_break(self):
        self.pastikan_seksi()
        p = self.doc.add_paragraph()
        p.add_run().add_break(WD_BREAK.PAGE)
        self.ada_isi = True

    def gambar(self, alt, path_rel):
        self.pastikan_seksi()
        path = (self.dasar / path_rel).resolve()
        if not path.exists():
            self.peringatan.append(f"Gambar tidak ditemukan: {path_rel}")
            self.paragraf(f"[GAMBAR TIDAK DITEMUKAN: {path_rel}]")
        else:
            shape = self.doc.add_picture(str(path))
            maks = Cm(F["page_width_cm"] - F["margin_left_cm"] - F["margin_right_cm"])
            if shape.width > maks:
                rasio = maks / shape.width
                shape.width = maks
                shape.height = int(shape.height * rasio)
            p = self.doc.paragraphs[-1]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.keep_with_next = True
        if alt.strip():
            cap = self.doc.add_paragraph()
            cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            cap.paragraph_format.line_spacing = 1.0
            cap.paragraph_format.space_after = Pt(6)
            add_inline(cap, alt.strip())
        self.ada_isi = True

    def tabel(self, baris_md):
        self.pastikan_seksi()
        rows = []
        for b in baris_md:
            sel = [c.strip() for c in b.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in sel if c):
                continue
            rows.append(sel)
        if not rows:
            return
        ncol = max(len(r) for r in rows)
        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = True
        tblpr = t._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for sisi, nilai in (("top", "single"), ("left", "nil"), ("bottom", "single"),
                            ("right", "nil"), ("insideH", "nil"), ("insideV", "nil")):
            e = OxmlElement(f"w:{sisi}")
            e.set(qn("w:val"), nilai)
            if nilai == "single":
                e.set(qn("w:sz"), "8")
                e.set(qn("w:space"), "0")
                e.set(qn("w:color"), "000000")
            borders.append(e)
        look = tblpr.find(qn("w:tblLook"))
        if look is not None:
            look.addprevious(borders)
        else:
            tblpr.append(borders)
        ukuran = F["table_font_pt"]
        for i, r in enumerate(rows):
            for j in range(ncol):
                teks = r[j] if j < len(r) else ""
                cell = t.cell(i, j)
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline(p, teks, size=ukuran, bold=True if i == 0 else None)
                if i == 0:
                    tcpr = cell._tc.get_or_add_tcPr()
                    tb = OxmlElement("w:tcBorders")
                    bb = OxmlElement("w:bottom")
                    bb.set(qn("w:val"), "single")
                    bb.set(qn("w:sz"), "8")
                    bb.set(qn("w:space"), "0")
                    bb.set(qn("w:color"), "000000")
                    tb.append(bb)
                    tcpr.append(tb)
        # lebar kolom: kolom pertama sempit bila berisi nomor ("No", "No.")
        total = F["page_width_cm"] - F["margin_left_cm"] - F["margin_right_cm"]
        t.autofit = False
        if ncol > 1 and re.match(r"^no\.?$", rows[0][0].strip().lower()):
            lebar = [1.2] + [(total - 1.2) / (ncol - 1)] * (ncol - 1)
        else:
            lebar = [total / ncol] * ncol
        for j, w in enumerate(lebar):
            t.columns[j].width = Cm(w)
            for cell in t.columns[j].cells:
                cell.width = Cm(w)
        # beri jarak kecil setelah tabel
        sp = self.doc.add_paragraph()
        sp.paragraph_format.line_spacing = 1.0
        sp.paragraph_format.space_after = Pt(0)
        self.ada_isi = True


# ---------------------------------------------------------------------------
def konversi(md_path: Path, tahap: str, romawi: bool, judul: str, penulis: str):
    doc = Document()
    cfg = config.TAHAP[tahap]
    setup_styles(doc, cfg["spasi"])
    S = Penyusun(doc, tahap, romawi, md_path.parent)

    baris = md_path.read_text(encoding="utf-8").splitlines()
    buf = []

    def flush():
        if buf:
            S.paragraf(" ".join(buf))
            buf.clear()

    i = 0
    while i < len(baris):
        s = baris[i].strip()
        if s.startswith("<!--"):
            if "pagebreak" in s.lower():
                flush()
                S.page_break()
            while "-->" not in baris[i]:
                i += 1
                if i >= len(baris):
                    break
            i += 1
            continue
        if s.startswith("```"):
            flush()
            blok = []
            i += 1
            while i < len(baris) and not baris[i].strip().startswith("```"):
                blok.append(baris[i].rstrip("\n"))
                i += 1
            i += 1
            S.kode(blok)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            flush()
            S.heading(m.group(2).strip(), len(m.group(1)))
            i += 1
            continue
        if s.startswith("|"):
            flush()
            blok = []
            while i < len(baris) and baris[i].strip().startswith("|"):
                blok.append(baris[i])
                i += 1
            S.tabel(blok)
            continue
        mi = IMG_RE.match(s)
        if mi:
            flush()
            S.gambar(mi.group(1), mi.group(2))
            i += 1
            continue
        mb = BULLET_RE.match(s)
        if mb and not s.startswith("**"):
            flush()
            S.butir_bullet(mb.group(1))
            i += 1
            continue
        mn = NUM_RE.match(s)
        if mn:
            flush()
            S.butir_nomor(mn.group(1), mn.group(2))
            i += 1
            continue
        if s == "":
            flush()
        else:
            buf.append(s)
        i += 1
    flush()

    cp = doc.core_properties
    cp.title = judul or md_path.stem
    cp.author = penulis
    cp.last_modified_by = penulis
    cp.comments = ""
    return doc, S.peringatan


def ke_pdf(docx_path: Path):
    """Coba LibreOffice dulu, lalu docx2pdf (butuh Microsoft Word di Windows/Mac)."""
    outdir = docx_path.parent
    kandidat = [shutil.which("soffice"), shutil.which("libreoffice"),
                r"C:\Program Files\LibreOffice\program\soffice.exe",
                r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                "/Applications/LibreOffice.app/Contents/MacOS/soffice"]
    for exe in kandidat:
        if exe and Path(exe).exists():
            try:
                subprocess.run([exe, "--headless", "--convert-to", "pdf", "--outdir", str(outdir),
                                str(docx_path)], check=True, timeout=240,
                               stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
                pdf = docx_path.with_suffix(".pdf")
                if pdf.exists():
                    return pdf
            except Exception as e:  # noqa: BLE001
                print(f"  LibreOffice gagal: {e}")
    try:
        from docx2pdf import convert  # type: ignore
        pdf = docx_path.with_suffix(".pdf")
        convert(str(docx_path), str(pdf))
        if pdf.exists():
            return pdf
    except ImportError:
        pass
    except Exception as e:  # noqa: BLE001
        print(f"  docx2pdf gagal: {e}")
    return None


def main():
    setup_stdout()
    ap = argparse.ArgumentParser(description="Markdown → .docx berformat pedoman TA FTI ITB Widya Gama.")
    ap.add_argument("berkas", help="file .md sumber")
    ap.add_argument("--tahap", choices=list(config.TAHAP), default="proposal")
    ap.add_argument("--keluar", help="path .docx keluaran (default: output/<nama>.docx)")
    ap.add_argument("--pdf", action="store_true", help="sekalian buat PDF (butuh LibreOffice atau Word)")
    ap.add_argument("--romawi", action="store_true",
                    help="nomor halaman Bab 1-5 angka Romawi (pembacaan harfiah Bab 5.1.B; default angka Arab)")
    ap.add_argument("--judul", default="", help="judul untuk properti dokumen")
    ap.add_argument("--penulis", default="", help="nama penulis untuk properti dokumen")
    a = ap.parse_args()

    md = Path(a.berkas)
    if not md.exists():
        print(f"Berkas tidak ditemukan: {md}")
        sys.exit(2)
    out = Path(a.keluar) if a.keluar else ensure_output_dir() / f"{md.stem}.docx"
    out.parent.mkdir(parents=True, exist_ok=True)

    doc, peringatan = konversi(md, a.tahap, a.romawi, a.judul, a.penulis)
    doc.save(str(out))
    print(f"✅ Dokumen dibuat: {out}")
    for w in peringatan:
        print(f"  ⚠️ {w}")

    if a.pdf:
        pdf = ke_pdf(out)
        if pdf:
            print(f"✅ PDF dibuat: {pdf}")
        else:
            print("⚠️ PDF tidak dibuat: LibreOffice tidak ditemukan dan docx2pdf tidak tersedia. "
                  "Buka .docx di Word lalu File → Save As → PDF.")
    print("Langkah berikut: python scripts/cek_format.py", out.as_posix())


if __name__ == "__main__":
    main()
